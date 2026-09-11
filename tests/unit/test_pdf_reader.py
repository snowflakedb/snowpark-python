#
# Copyright (c) 2012-2025 Snowflake Computing Inc. All rights reserved.
#

import io
import json
import os
from unittest.mock import patch

import pytest

from snowflake.snowpark._internal.pdf_reader import (
    PAGE_SEPARATOR,
    PDFTextReader,
    extract_docx_pages,
    extract_pdf_pages,
    extract_text_pages,
    selected_page_indexes,
)
from tests.utils import TestFiles

RESOURCES_PATH = os.path.normpath(
    os.path.join(os.path.dirname(__file__), "../resources")
)
TEST_FILES = TestFiles(RESOURCES_PATH)

DOC_PAGE_COUNT = 3


def _page_filter(*ranges):
    return json.dumps([{"start": start, "end": end} for start, end in ranges])


def test_selected_page_indexes_empty_filter_selects_every_page():
    assert selected_page_indexes(5, "") == [0, 1, 2, 3, 4]


def test_selected_page_indexes_single_range():
    assert selected_page_indexes(5, _page_filter((1, 3))) == [1, 2]


def test_selected_page_indexes_multiple_ranges_are_merged_and_sorted():
    assert selected_page_indexes(5, _page_filter((3, 5), (0, 2))) == [0, 1, 3, 4]


def test_selected_page_indexes_overlapping_ranges_deduplicate():
    assert selected_page_indexes(5, _page_filter((0, 3), (1, 4))) == [0, 1, 2, 3]


def test_selected_page_indexes_clamps_out_of_bounds_range():
    assert selected_page_indexes(3, _page_filter((-2, 10))) == [0, 1, 2]


def test_selected_page_indexes_empty_range_selects_nothing():
    assert selected_page_indexes(5, _page_filter((2, 2))) == []


# ---------------------------------------------------------------------------
# extract_pdf_pages
# ---------------------------------------------------------------------------


def test_extract_pdf_pages_returns_every_page_by_default():
    with open(TEST_FILES.test_doc_pdf, "rb") as f:
        total_pages, pages = extract_pdf_pages(f, "")

    assert total_pages == DOC_PAGE_COUNT
    assert [index for index, _ in pages] == [0, 1, 2]
    assert all(text.strip() for _, text in pages)


def test_extract_pdf_pages_honors_page_filter():
    with open(TEST_FILES.test_doc_pdf, "rb") as f:
        total_pages, pages = extract_pdf_pages(f, _page_filter((0, 2)))

    assert total_pages == DOC_PAGE_COUNT
    assert [index for index, _ in pages] == [0, 1]


def test_extract_pdf_pages_empty_selection_returns_no_pages():
    with open(TEST_FILES.test_doc_pdf, "rb") as f:
        total_pages, pages = extract_pdf_pages(f, _page_filter((2, 2)))

    assert total_pages == DOC_PAGE_COUNT
    assert pages == []


# ---------------------------------------------------------------------------
# extract_docx_pages
# ---------------------------------------------------------------------------


def _build_docx_bytes(sections):
    """Build a minimal in-memory DOCX with one Heading + body paragraph per
    section, so extract_docx_pages' heading-boundary splitting has something
    real to split on."""
    import docx

    document = docx.Document()
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf


def test_extract_docx_pages_splits_on_headings():
    buf = _build_docx_bytes(
        [("Section One", "First body text"), ("Section Two", "Second body text")]
    )
    total_pages, pages = extract_docx_pages(buf, "")

    assert total_pages == 2
    assert "Section One" in pages[0][1]
    assert "First body text" in pages[0][1]
    assert "Section Two" in pages[1][1]
    assert "Second body text" in pages[1][1]


def test_extract_docx_pages_no_headings_is_a_single_page():
    import docx

    document = docx.Document()
    document.add_paragraph("Just one plain paragraph, no heading.")
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    total_pages, pages = extract_docx_pages(buf, "")
    assert total_pages == 1
    assert "Just one plain paragraph" in pages[0][1]


def test_extract_docx_pages_skips_blank_paragraphs():
    import docx

    document = docx.Document()
    document.add_heading("Section One", level=1)
    document.add_paragraph("First body text")
    document.add_paragraph("   ")  # blank paragraph, should be skipped
    document.add_paragraph("More text after the blank line")
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    total_pages, pages = extract_docx_pages(buf, "")
    assert total_pages == 1
    assert "First body text" in pages[0][1]
    assert "More text after the blank line" in pages[0][1]


def test_extract_docx_pages_all_blank_is_zero_pages():
    import docx

    document = docx.Document()
    document.add_paragraph("   ")
    document.add_paragraph("")
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    total_pages, pages = extract_docx_pages(buf, "")
    assert total_pages == 0
    assert pages == []


def test_extract_docx_pages_honors_page_filter():
    buf = _build_docx_bytes(
        [
            ("Section One", "First"),
            ("Section Two", "Second"),
            ("Section Three", "Third"),
        ]
    )
    total_pages, pages = extract_docx_pages(buf, _page_filter((0, 2)))

    assert total_pages == 3
    assert [index for index, _ in pages] == [0, 1]


# ---------------------------------------------------------------------------
# extract_text_pages
# ---------------------------------------------------------------------------


def test_extract_text_pages_short_file_is_a_single_page():
    buf = io.BytesIO(b"line one\nline two\nline three\n")
    total_pages, pages = extract_text_pages(buf, "")

    assert total_pages == 1
    assert pages[0][0] == 0
    assert "line one" in pages[0][1]
    assert "line three" in pages[0][1]


def test_extract_text_pages_splits_every_100_lines():
    content = "".join(f"line {i}\n" for i in range(250))
    buf = io.BytesIO(content.encode("utf-8"))
    total_pages, pages = extract_text_pages(buf, "")

    assert total_pages == 3  # 100 + 100 + 50
    assert "line 0" in pages[0][1]
    assert "line 150" in pages[1][1]
    assert "line 249" in pages[2][1]


def test_extract_text_pages_honors_page_filter():
    content = "".join(f"line {i}\n" for i in range(250))
    buf = io.BytesIO(content.encode("utf-8"))
    total_pages, pages = extract_text_pages(buf, _page_filter((1, 3)))

    assert total_pages == 3
    assert [index for index, _ in pages] == [1, 2]


# ---------------------------------------------------------------------------
# PDFTextReader.process()
# ---------------------------------------------------------------------------


def _mock_snowflake_file_open(fileobj_factory):
    """SnowflakeFile.open() is only meaningful inside a real UDTF; the rest of
    PDFTextReader.process() is plain Python and doesn't need a live session, as
    long as SnowflakeFile.open() is replaced with a local/in-memory file handle."""
    return patch(
        "snowflake.snowpark._internal.pdf_reader.SnowflakeFile.open",
        side_effect=lambda *a, **k: fileobj_factory(),
    )


def test_process_document_boundary_joins_pages():
    with _mock_snowflake_file_open(lambda: open(TEST_FILES.test_doc_pdf, "rb")):
        rows = list(
            PDFTextReader().process("@stage/doc.pdf", "document", "", "PERMISSIVE")
        )

    assert len(rows) == 1
    index, total_pages, content, error = rows[0]
    assert index is None
    assert total_pages == DOC_PAGE_COUNT
    assert error is None
    assert content.count(PAGE_SEPARATOR) == DOC_PAGE_COUNT - 1


def test_process_page_boundary_yields_one_row_per_page():
    with _mock_snowflake_file_open(lambda: open(TEST_FILES.test_doc_pdf, "rb")):
        rows = list(PDFTextReader().process("@stage/doc.pdf", "page", "", "PERMISSIVE"))

    assert [r[0] for r in rows] == [0, 1, 2]
    assert all(r[1] == DOC_PAGE_COUNT for r in rows)
    assert all(r[3] is None for r in rows)


def test_process_page_boundary_honors_page_filter():
    with _mock_snowflake_file_open(lambda: open(TEST_FILES.test_doc_pdf, "rb")):
        rows = list(
            PDFTextReader().process(
                "@stage/doc.pdf", "page", _page_filter((0, 2)), "PERMISSIVE"
            )
        )

    assert [r[0] for r in rows] == [0, 1]


def test_process_docx_document_boundary():
    with _mock_snowflake_file_open(
        lambda: _build_docx_bytes([("Section One", "First"), ("Section Two", "Second")])
    ):
        rows = list(
            PDFTextReader().process("@stage/doc.docx", "document", "", "PERMISSIVE")
        )

    assert len(rows) == 1
    index, total_pages, content, error = rows[0]
    assert index is None
    assert total_pages == 2
    assert error is None
    assert "First" in content and "Second" in content


def test_process_docx_page_boundary():
    with _mock_snowflake_file_open(
        lambda: _build_docx_bytes([("Section One", "First"), ("Section Two", "Second")])
    ):
        rows = list(
            PDFTextReader().process("@stage/doc.docx", "page", "", "PERMISSIVE")
        )

    assert [r[0] for r in rows] == [0, 1]
    assert all(r[1] == 2 for r in rows)


def test_process_txt_document_boundary():
    with _mock_snowflake_file_open(
        lambda: io.BytesIO(b"hello from a plain text transcript\n")
    ):
        rows = list(
            PDFTextReader().process(
                "@stage/transcript.txt", "document", "", "PERMISSIVE"
            )
        )

    assert len(rows) == 1
    index, total_pages, content, error = rows[0]
    assert index is None
    assert total_pages == 1
    assert error is None
    assert "hello from a plain text transcript" in content


def test_process_md_file_uses_text_extraction():
    with _mock_snowflake_file_open(
        lambda: io.BytesIO(b"# Title\n\nSome markdown body.\n")
    ):
        rows = list(
            PDFTextReader().process("@stage/notes.md", "document", "", "PERMISSIVE")
        )

    assert len(rows) == 1
    assert "Some markdown body." in rows[0][2]


def test_process_unsupported_extension_is_permissive_error():
    with _mock_snowflake_file_open(lambda: open(TEST_FILES.test_audio_ogg, "rb")):
        rows = list(
            PDFTextReader().process("@stage/audio.ogg", "document", "", "PERMISSIVE")
        )

    assert len(rows) == 1
    index, total_pages, content, error = rows[0]
    assert index is None
    assert total_pages is None
    assert content is None
    assert "audio.ogg" in error
    assert "does not support file type" in error


def test_process_unsupported_extension_failfast_raises():
    with _mock_snowflake_file_open(lambda: open(TEST_FILES.test_audio_ogg, "rb")):
        with pytest.raises(RuntimeError, match="does not support file type"):
            list(
                PDFTextReader().process("@stage/audio.ogg", "document", "", "FAILFAST")
            )


def test_process_permissive_yields_an_error_row_for_an_unreadable_pdf():
    with _mock_snowflake_file_open(lambda: open(TEST_FILES.test_audio_ogg, "rb")):
        rows = list(
            PDFTextReader().process("@stage/audio.pdf", "document", "", "PERMISSIVE")
        )

    assert len(rows) == 1
    index, total_pages, content, error = rows[0]
    assert index is None
    assert total_pages is None
    assert content is None
    assert "audio.pdf" in error


def test_process_failfast_raises_for_an_unreadable_pdf():
    with _mock_snowflake_file_open(lambda: open(TEST_FILES.test_audio_ogg, "rb")):
        with pytest.raises(RuntimeError, match="Failed to extract text"):
            list(
                PDFTextReader().process("@stage/audio.pdf", "document", "", "FAILFAST")
            )
